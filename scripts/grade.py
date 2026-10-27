from openpyxl import load_workbook
import sys
f,lo,hi=sys.argv[1],float(sys.argv[2]),float(sys.argv[3])
ws=load_workbook(f,data_only=True).active
rows=[r for r in ws.iter_rows(min_row=3,values_only=True) if r[0]]
print(ws['A1'].value, 'spot~',round(rows[0][6]/(1+rows[0][8]),2))
for r in rows:
  if r[1].month==12 and r[1].year==2026 and lo<=r[6]<=hi:
    print(r[0].strip(),r[3][0],r[6],'ult',r[9],str(r[11])[-8:],'neg',r[12],'IV',r[14])
