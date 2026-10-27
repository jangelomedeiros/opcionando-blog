from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

F='Arial'
f=lambda **k: Font(name=F,size=k.pop('size',10),**k)
AMAR=PatternFill('solid',fgColor='FFFF00'); CAB=PatternFill('solid',fgColor='0E4A55')
CINZA=PatternFill('solid',fgColor='EEF2F3'); CORAL=PatternFill('solid',fgColor='FCE3DC')
fino=Side(style='thin',color='C9D3D6'); B=Border(left=fino,right=fino,top=fino,bottom=fino)
AZUL='0000FF'; VERDE='008000'

wb=Workbook()
# ---------------- Leia-me
lm=wb.active; lm.title='Leia-me'
txt=[('Opcionando · Cotações das estruturas via RTD do Profit',dict(bold=True,size=14,color='0E4A55')),
('',{}),
('Como usar',dict(bold=True,size=11)),
('1. Abra o Profit e mantenha-o conectado. Abra esta planilha no Excel (no Windows, onde o Profit está instalado).',{}),
('2. Na aba "Opções", preencha só as células AMARELAS da coluna "Código da série" com o código de cada opção (ex.: VALEL804).',{}),
('   A coluna "Começa com" mostra a raiz e a letra de dezembro: calls de dezembro usam L e puts de dezembro usam X.',{}),
('3. Confira a coluna "Strike confere?": se aparecer "CONFIRA", o código digitado não tem o strike da análise.',{}),
('4. Os campos de oferta de compra e de venda estão em branco na linha 4 da aba "Opções" (células amarelas).',{}),
('   Copie do Profit o código RTD desses dois campos e cole nelas. Os demais códigos (ULT, PEX, NEG...) vieram da sua lista.',{}),
('5. A aba "Estruturas" recalcula crédito, débito, perda máxima, ganho máximo e equilíbrio com os prêmios reais.',{}),
('   Enquanto não houver cotação, ela usa o prêmio teórico (Black-Scholes) da análise e indica isso na coluna "Fonte".',{}),
('6. Salve o arquivo (os últimos valores ficam gravados) e me envie. Eu atualizo as análises do site.',{}),
('',{}),
('Legenda de cores',dict(bold=True,size=11)),
('Amarelo: célula para você preencher.',{}),
('Azul: valor fixo vindo das análises de 26/10 (strike e prêmio teórico).',{}),
('Preto: fórmula. Verde: fórmula que puxa dados de outra aba.',{}),
('',{}),
('Observação: as fórmulas RTD só retornam valores no Excel do Windows com o Profit aberto. Em outro programa, as células aparecem vazias ou com erro.',dict(italic=True,color='666666')),
]
for i,(t,st) in enumerate(txt,1):
    lm.cell(i,1,t).font=f(**st)
lm.column_dimensions['A'].width=120

# ---------------- Opções
op=wb.create_sheet('Opções')
op['A1']='Séries das estruturas · vencimento 18/12/2026'; op['A1'].font=f(bold=True,size=13,color='0E4A55')
op['A2']='Sufixo RTD das opções e ações à vista:'; op['A2'].font=f(italic=True)
op['E2']='_B_0'; op['E2'].font=f(color=AZUL,bold=True); op['E2'].fill=AMAR
op['E2'].comment=Comment('Sufixo do mercado Bovespa no RTD do Profit, conforme a lista enviada pelo Ângelo (ex.: VALEL804_B_0).','Claude')
op['A4']='Código do campo RTD →'; op['A4'].font=f(italic=True,color='666666')

fixas=['Análise','Perna','Tipo','Strike (análise)','Prêmio teórico','Começa com','Código da série']
rtd=[('Data','DAT'),('Hora','HOR'),('Último','ULT'),('Abertura','ABE'),('Máximo','MAX'),('Mínimo','MIN'),
     ('Strike (PEX)','PEX'),('Negócios','NEG'),('Vencimento','VEN'),('Oferta compra',''),('Oferta venda','')]
calc=['Strike confere?','Prêmio usado','Fonte','Sinal','Strike usado','Real − teórico']
cab=fixas+[r[0] for r in rtd]+calc
HR=5
for j,h in enumerate(cab,1):
    c=op.cell(HR,j,h); c.font=f(bold=True,color='FFFFFF'); c.fill=CAB; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=B
op.row_dimensions[HR].height=30
c0=len(fixas)+1  # primeira coluna RTD
col={}
for j,(nome,cod) in enumerate(rtd):
    cc=c0+j; col[nome]=L(cc)
    c=op.cell(4,cc,cod); c.font=f(color=AZUL,bold=True); c.alignment=Alignment(horizontal='center')
    if not cod:
        c.fill=AMAR; c.comment=Comment('Cole aqui o código RTD deste campo, copiado do Profit.','Claude')
k=c0+len(rtd)
cC={n:L(k+i) for i,n in enumerate(calc)}

pernas=[
 ('VALE3','Compra put','put',56,0.46,'VALEX'),
 ('VALE3','Venda put','put',60,1.23,'VALEX'),
 ('VALE3','Venda call','call',70,1.79,'VALEL'),
 ('VALE3','Compra call','call',74,0.91,'VALEL'),
 ('PETR4','Venda put','put',50,1.96,'PETRX'),
 ('PETR4','Compra put','put',46,0.89,'PETRX'),
 ('BBAS3','Venda call','call',20,1.59,'BBASL'),
 ('BBAS3','Compra call','call',22,1.08,'BBASL'),
 ('BBDC4','Venda put','put',14,0.90,'BBDCX'),
 ('PRIO3','Compra call','call',72,4.02,'PRIOL'),
 ('PRIO3','Venda call','call',78,1.80,'PRIOL'),
]
r0=HR+1
for i,p in enumerate(pernas):
    r=r0+i
    vals=list(p[:5])+[p[5]+'…']
    for j,v in enumerate(vals,1):
        c=op.cell(r,j,v); c.border=B; c.font=f(color=AZUL if j in (4,5) else '000000')
    op.cell(r,4).number_format='0.00'; op.cell(r,5).number_format='0.00'
    g=op.cell(r,7); g.fill=AMAR; g.border=B; g.font=f(bold=True)
    for j,(nome,cod) in enumerate(rtd):
        cl=L(c0+j)
        c=op.cell(r,c0+j,f'=IF(OR($G{r}="",{cl}$4=""),"",IFERROR(RTD("RTDTrading.RTDServer",,$G{r}&$E$2,{cl}$4),""))')
        c.border=B; c.font=f()
        if nome in ('Último','Abertura','Máximo','Mínimo','Strike (PEX)','Oferta compra','Oferta venda'): c.number_format='0.00'
    ult,bid,ask,pex=col['Último'],col['Oferta compra'],col['Oferta venda'],col['Strike (PEX)']
    F_=lambda n,v: (op.cell(r,list(cC).index(n)+k,v))
    F_('Strike confere?',f'=IF(ISNUMBER({pex}{r}),IF(ABS({pex}{r}-D{r})<0.005,"ok","CONFIRA"),"")')
    F_('Prêmio usado',f'=IF(AND(ISNUMBER({bid}{r}),ISNUMBER({ask}{r}),{bid}{r}>0,{ask}{r}>0),({bid}{r}+{ask}{r})/2,IF(ISNUMBER({ult}{r}),{ult}{r},E{r}))')
    F_('Fonte',f'=IF(AND(ISNUMBER({bid}{r}),ISNUMBER({ask}{r}),{bid}{r}>0,{ask}{r}>0),"meio do book",IF(ISNUMBER({ult}{r}),"último negócio","teórico"))')
    F_('Sinal',f'=IF(LEFT(B{r},5)="Venda",1,-1)')
    F_('Strike usado',f'=IF(ISNUMBER({pex}{r}),{pex}{r},D{r})')
    F_('Real − teórico',f'={cC["Prêmio usado"]}{r}-E{r}')
    for n in calc:
        c=op[f'{cC[n]}{r}']; c.border=B; c.font=f()
        if n in ('Prêmio usado','Strike usado','Real − teórico'): c.number_format='0.00;-0.00;-'
    if i%2: 
        for j in range(1,len(cab)+1):
            if op.cell(r,j).fill!=AMAR: op.cell(r,j).fill=CINZA
rN=r0+len(pernas)-1
op['A'+str(rN+2)]='Prêmio usado: meio do book (média das ofertas de compra e de venda) quando houver ofertas; senão o último negócio; senão o prêmio teórico da análise.'
op['A'+str(rN+2)].font=f(italic=True,color='666666')
op['A'+str(rN+3)]='Strike e prêmio teórico: análises de 26/10/2026 do opcionando.com.br (Black-Scholes, Selic 14% a.a., vencimento 18/12/2026).'
op['A'+str(rN+3)].font=f(italic=True,color='666666')

# ativos à vista
ra=rN+6
op.cell(ra-1,1,'Ativos-objeto').font=f(bold=True,size=11,color='0E4A55')
for j,h in enumerate(['Ativo','Último','Máximo','Mínimo','Negócios','Data','Hora'],1):
    c=op.cell(ra,j,h); c.font=f(bold=True,color='FFFFFF'); c.fill=CAB; c.border=B; c.alignment=Alignment(horizontal='center')
cods=['ULT','MAX','MIN','NEG','DAT','HOR']
for i,t in enumerate(['VALE3','PETR4','BBAS3','BBDC4','PRIO3']):
    r=ra+1+i; op.cell(r,1,t).font=f(bold=True); op.cell(r,1).border=B
    for j,cd in enumerate(cods,2):
        c=op.cell(r,j,f'=IFERROR(RTD("RTDTrading.RTDServer",,$A{r}&$E$2,"{cd}"),"")'); c.border=B; c.font=f()
        if cd in('ULT','MAX','MIN'): c.number_format='0.00'
SP={t:f'$B${ra+1+i}' for i,t in enumerate(['VALE3','PETR4','BBAS3','BBDC4','PRIO3'])}

widths={'A':9,'B':12,'C':6,'D':9,'E':9,'F':10,'G':13}
for c_,w in widths.items(): op.column_dimensions[c_].width=w
for j in range(c0,len(cab)+1): op.column_dimensions[L(j)].width=11
op.freeze_panes='H6'

# ---------------- Estruturas
es=wb.create_sheet('Estruturas')
es['A1']='Estruturas recalculadas com os prêmios reais'; es['A1'].font=f(bold=True,size=13,color='0E4A55')
es['A2']='Valores por lote de 100 opções. Positivo = crédito recebido; negativo = débito pago.'; es['A2'].font=f(italic=True,color='666666')
hd=['Análise','Estrutura','Montagem real (R$/ação)','Montagem teórica (R$/ação)','Diferença','Montagem real (R$/lote)','Ganho máximo (R$/lote)','Perda máxima (R$/lote)','Equilíbrio inferior','Equilíbrio superior','Fonte dos prêmios','Preço do ativo']
for j,h in enumerate(hd,1):
    c=es.cell(4,j,h); c.font=f(bold=True,color='FFFFFF'); c.fill=CAB; c.border=B; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
es.row_dimensions[4].height=42
O="'Opções'!"
rr=lambda a: f"{O}${a}${r0}:${a}${rN}"
PU,SI,SU=cC['Prêmio usado'],cC['Sinal'],cC['Strike usado']
def linha(t): return [i for i,p in enumerate(pernas) if p[0]==t]
def su(t,idx): return f"{O}${SU}${r0+linha(t)[idx]}"
estr=[('VALE3','Iron condor 56/60/70/74'),('PETR4','Trava de alta com put 50/46'),('BBAS3','Trava de baixa com call 20/22'),('BBDC4','Venda de put coberta 14'),('PRIO3','Trava de alta com call 72/78')]
for i,(t,n) in enumerate(estr):
    r=5+i
    es.cell(r,1,t).font=f(bold=True); es.cell(r,2,n).font=f()
    es.cell(r,3,f'=SUMPRODUCT(({rr("A")}=$A{r})*{rr(SI)}*{rr(PU)})').font=f(color=VERDE)
    es.cell(r,4,f'=SUMPRODUCT(({rr("A")}=$A{r})*{rr(SI)}*{rr("E")})').font=f(color=VERDE)
    es.cell(r,5,f'=C{r}-D{r}')
    es.cell(r,6,f'=C{r}*100')
    if t=='VALE3':
        lp,sp_,sc,lc=(su(t,x) for x in range(4))
        g=f'=C{r}*100'; pm=f'=-(MAX({sp_}-{lp},{lc}-{sc})-C{r})*100'; e1=f'={sp_}-C{r}'; e2=f'={sc}+C{r}'
    elif t=='PETR4':
        s,b_=su(t,0),su(t,1); g=f'=C{r}*100'; pm=f'=-(({s}-{b_})-C{r})*100'; e1=f'={s}-C{r}'; e2='=""'
    elif t=='BBAS3':
        s,b_=su(t,0),su(t,1); g=f'=C{r}*100'; pm=f'=-(({b_}-{s})-C{r})*100'; e1='=""'; e2=f'={s}+C{r}'
    elif t=='BBDC4':
        s=su(t,0); g=f'=C{r}*100'; pm=f'=-({s}-C{r})*100'; e1=f'={s}-C{r}'; e2='=""'
    else:
        b_,s=su(t,0),su(t,1); g=f'=(({s}-{b_})+C{r})*100'; pm=f'=C{r}*100'; e1=f'={b_}-C{r}'; e2='=""'
    for cc,v in zip('GHIJ',[g,pm,e1,e2]): es[f'{cc}{r}']=v
    idx=linha(t)
    fontes=','.join(f'{O}{cC["Fonte"]}{r0+x}' for x in idx)
    es.cell(r,11,f'=IF(COUNTIF({O}${cC["Fonte"]}${r0+idx[0]}:${cC["Fonte"]}${r0+idx[-1]},"teórico")={len(idx)},"teórico",IF(COUNTIF({O}${cC["Fonte"]}${r0+idx[0]}:${cC["Fonte"]}${r0+idx[-1]},"teórico")=0,"real","parcial"))').font=f(color=VERDE)
    es.cell(r,12,f'=IFERROR({O}{SP[t]}*1,"")').font=f(color=VERDE)
    for j in range(1,13):
        c=es.cell(r,j); c.border=B
        if c.font.name!=F: c.font=f()
        if j in(3,4,5,9,10,12): c.number_format='0.00;-0.00;-'
        if j in(6,7,8): c.number_format='#,##0.00;-#,##0.00;-'
    if i%2:
        for j in range(1,13): es.cell(r,j).fill=CINZA
es['A11']='Perda máxima da put coberta (BBDC4): caixa comprometido se a ação for a zero, isto é, (strike − prêmio) × 100.'
es['A12']='Fonte "teórico": nenhuma perna com cotação ainda; "parcial": parte das pernas com cotação real; "real": todas com cotação.'
for a in('A11','A12'): es[a].font=f(italic=True,color='666666')
for c_,w in zip('ABCDEFGHIJKL',[9,30,13,13,11,13,13,13,11,11,11,10]): es.column_dimensions[c_].width=w
es.freeze_panes='C5'

wb.calculation.fullCalcOnLoad=True
wb.move_sheet('Opções',-1)  # Opções primeiro
wb.active=0
wb.save('opcionando-rtd-estruturas.xlsx')
print(r0,rN,cC)
